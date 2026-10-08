"""Turn pasted text or an uploaded document into retrievable chunks."""

import io
import re
import unicodedata

SUPPORTED_EXTENSIONS = (".txt", ".md", ".docx", ".pdf")

MAX_CHUNK_CHARS = 1100
MIN_CHUNK_CHARS = 120
OVERLAP_CHARS = 150


class ExtractionError(ValueError):
    pass


def extract_text(filename, blob):
    """Pull plain text out of an uploaded file."""
    name = (filename or "").lower()

    if name.endswith(".docx"):
        return _from_docx(blob)
    if name.endswith(".pdf"):
        return _from_pdf(blob)
    if name.endswith((".txt", ".md")):
        return blob.decode("utf-8", errors="replace")
    if name.endswith(".doc"):
        raise ExtractionError(
            "Legacy .doc isn't supported — re-save it as .docx and try again."
        )
    raise ExtractionError(
        "Unsupported file type. Use one of: " + ", ".join(SUPPORTED_EXTENSIONS)
    )


def _from_docx(blob):
    import docx

    document = docx.Document(io.BytesIO(blob))
    parts = [para.text for para in document.paragraphs]

    # Tables often hold the actual specs in a client brief; don't drop them.
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    return "\n".join(parts)


def _from_pdf(blob):
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(blob))
    return "\n\n".join((page.extract_text() or "") for page in reader.pages)


def clean(text):
    text = unicodedata.normalize("NFKC", text or "")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _split_sentences(paragraph):
    return re.split(r"(?<=[.!?])\s+", paragraph)


def chunk(text, max_chars=MAX_CHUNK_CHARS, overlap=OVERLAP_CHARS):
    """Split on paragraph boundaries, falling back to sentences.

    Chunks carry a tail of the previous chunk so a fact split across a
    boundary is still retrievable from either side.
    """
    text = clean(text)
    if not text:
        return []

    pieces = []
    for paragraph in text.split("\n\n"):
        paragraph = paragraph.strip()
        if not paragraph:
            continue
        if len(paragraph) <= max_chars:
            pieces.append(paragraph)
            continue

        buffer = ""
        for sentence in _split_sentences(paragraph):
            if len(buffer) + len(sentence) + 1 > max_chars and buffer:
                pieces.append(buffer.strip())
                buffer = sentence
            else:
                buffer = (buffer + " " + sentence).strip()
        if buffer:
            pieces.append(buffer.strip())

    # Merge runs of short paragraphs, then add overlap.
    merged = []
    for piece in pieces:
        if merged and len(merged[-1]) + len(piece) + 2 <= max_chars:
            merged[-1] = merged[-1] + "\n\n" + piece
        else:
            merged.append(piece)

    chunks = []
    for index, piece in enumerate(merged):
        if index and overlap:
            tail = merged[index - 1][-overlap:]
            tail = tail[tail.find(" ") + 1 :] if " " in tail else tail
            piece = tail.strip() + " … " + piece
        chunks.append(piece)

    # A trailing scrap alone isn't worth a chunk; fold it back.
    if len(chunks) > 1 and len(chunks[-1]) < MIN_CHUNK_CHARS:
        chunks[-2] = chunks[-2] + "\n\n" + chunks[-1]
        chunks.pop()

    return chunks
